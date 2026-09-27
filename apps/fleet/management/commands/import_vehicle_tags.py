"""Set each vehicle's RFID tag and branch from the client's Excel, and make it active.

Source: the "Branch Vehicle RFTag" sheet of the client's OLD RMS Data.xlsx --
one row per vehicle: Branch, Category, VehicleType, Vehicle (its name), RFTagID.

Why this exists: the first import (import_live_vehicle_data) took rfid_epc from
imsstockitem.BarCode, a printed label number. The legacy system kept the real
tag in RmsVehicleRFTagDetails.TagID, which the client's sheet now supplies. It
also gives the branch each vehicle stands at; current_branch was empty for all.

Rules, decided with the user:
  - A row matches a vehicle of the company by name (case and spacing ignored)
    whose category and vehicle type are the ones the row names -- a name alone
    could land on the wrong vehicle. Several matches: the one active vehicle
    wins; more than one active is ambiguous.
  - Each matched vehicle gets rfid_epc (the tag as text, upper case),
    current_branch, and is_active=True. Nothing else changes, and vehicles not
    in the sheet are not touched.
  - Tags are text: Excel turns a tag made only of digits into a number, so it
    is read back as its digits. Every tag must be unique within the sheet.
  - Another vehicle still holding the same value in rfid_epc holds an old
    barcode, never a tag (see above): that value is cleared.
  - Anything unresolved -- an unknown vehicle, branch or type, an ambiguous
    name, a blank or repeated tag -- is listed and nothing is written.
  - Dry run unless --commit: the report shows what would change.
"""

import re
from collections import Counter, defaultdict

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.company.models import Branch, Company
from apps.fleet.models import Vehicle

SHEET = "Branch Vehicle RFTag"
COLUMNS = ("Branch", "Category", "VehicleType", "Vehicle", "RFTagID")


def norm(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def tag_text(value):
    """A tag as text. Excel stores 17348246 as a number and 170732EA as text."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip().upper()


class Command(BaseCommand):
    help = "Set RFID tag, branch and active on each vehicle in the client's Excel (dry run unless --commit)."

    def add_arguments(self, parser):
        parser.add_argument("xlsx", help="Path to the client's workbook")
        parser.add_argument("--company", required=True, help="Company short_code the vehicles belong to")
        parser.add_argument("--commit", action="store_true", help="Write the changes (default: dry run)")

    def handle(self, *args, **options):
        company = Company.objects.filter(short_code=options["company"]).first()
        if company is None:
            raise CommandError(f"No company with short_code={options['company']!r}.")
        rows = self.read(options["xlsx"])
        plan, problems = self.plan(company, rows)
        if problems:
            for problem in problems:
                self.stderr.write(f"  {problem}")
            raise CommandError(f"{len(problems)} row(s) could not be resolved; nothing was written.")

        with transaction.atomic():
            cleared, activated = self.apply(company, plan)
            report = self.report(plan, cleared, activated)
            if not options["commit"]:
                transaction.set_rollback(True)
        self.stdout.write(report)
        self.stdout.write(self.style.SUCCESS("Committed.") if options["commit"]
                          else self.style.WARNING("Dry run: nothing written. Re-run with --commit."))

    # -- Reading ----------------------------------------------------------------------

    def read(self, path):
        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except (OSError, ValueError) as error:
            raise CommandError(f"Cannot open {path}: {error}") from error
        if SHEET not in workbook.sheetnames:
            raise CommandError(f"No sheet named {SHEET!r} (found: {', '.join(workbook.sheetnames)}).")
        cells = workbook[SHEET].iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(cells, ())]
        missing = [c for c in COLUMNS if c not in header]
        if missing:
            raise CommandError(f"Sheet {SHEET!r} lacks column(s): {', '.join(missing)}.")
        index = {c: header.index(c) for c in COLUMNS}
        rows = []
        for number, values in enumerate(cells, start=2):
            if not any(v not in (None, "") for v in values):
                continue
            row = {c: values[index[c]] if index[c] < len(values) else None for c in COLUMNS}
            row["line"] = number
            rows.append(row)
        return rows

    # -- Matching ---------------------------------------------------------------------

    def plan(self, company, rows):
        branches = defaultdict(list)
        for branch in Branch.objects.filter(company=company):
            branches[norm(branch.name)].append(branch)
        vehicles = defaultdict(list)
        for vehicle in Vehicle.objects.filter(company=company).select_related("vehicle_type__category"):
            vehicles[norm(vehicle.vehicle_name)].append(vehicle)

        tags = Counter(tag_text(r["RFTagID"]) for r in rows)
        plan, problems, taken = [], [], set()
        for row in rows:
            where = f"line {row['line']} ({row['Vehicle']!s}):"
            tag = tag_text(row["RFTagID"])
            if not tag:
                problems.append(f"{where} no RFTagID")
                continue
            if tags[tag] > 1:
                problems.append(f"{where} tag {tag} appears {tags[tag]} times in the sheet")
                continue
            branch = branches.get(norm(row["Branch"]), [])
            if len(branch) != 1:
                problems.append(f"{where} branch {row['Branch']!r} "
                                f"{'not found' if not branch else 'matches several branches'}")
                continue
            named = vehicles.get(norm(row["Vehicle"]), [])
            if not named:
                problems.append(f"{where} no vehicle of that name")
                continue
            typed = [v for v in named if norm(v.vehicle_type.vehicle_type_name) == norm(row["VehicleType"])
                     and norm(v.vehicle_type.category.category_name) == norm(row["Category"])]
            if not typed:
                found = ", ".join(sorted({f"{v.vehicle_type.category.category_name}/{v.vehicle_type.vehicle_type_name}"
                                          for v in named}))
                problems.append(f"{where} is {found} here, not {row['Category']}/{row['VehicleType']}")
                continue
            active = [v for v in typed if v.is_active]
            chosen = active if active else typed
            if len(chosen) != 1:
                codes = ", ".join(v.vehicle_code for v in chosen)
                problems.append(f"{where} matches several vehicles ({codes})")
                continue
            vehicle = chosen[0]
            if vehicle.pk in taken:
                problems.append(f"{where} vehicle {vehicle.vehicle_code} is already on another line")
                continue
            taken.add(vehicle.pk)
            plan.append((vehicle, tag, branch[0]))
        return plan, problems

    # -- Writing ----------------------------------------------------------------------

    def apply(self, company, plan):
        """Clear old barcodes that equal a new tag, then write every vehicle.
        Returns (codes whose old value was cleared, how many were made active)."""
        mine = {vehicle.pk for vehicle, _, _ in plan}
        new_tags = {tag for _, tag, _ in plan}
        holders = Vehicle.objects.filter(company=company, rfid_epc__in=new_tags).exclude(pk__in=mine)
        cleared = sorted(holders.values_list("vehicle_code", flat=True))
        holders.update(rfid_epc="")
        # Blank first, so a tag moving between two listed vehicles never meets
        # itself under the unique constraint.
        Vehicle.objects.filter(pk__in=mine).update(rfid_epc="")
        activated = 0
        for vehicle, tag, branch in plan:
            activated += not vehicle.is_active
            vehicle.rfid_epc, vehicle.current_branch, vehicle.is_active = tag, branch, True
            vehicle.save(update_fields=["rfid_epc", "current_branch", "is_active", "modified_on"])
        return cleared, activated

    def report(self, plan, cleared, activated):
        per_branch = Counter(branch.name for _, _, branch in plan)
        lines = [
            f"{len(plan)} vehicles: RFID tag and branch set.",
            f"  made active: {activated} (the rest already were)",
            f"  across {len(per_branch)} branches: "
            + ", ".join(f"{name} {count}" for name, count in per_branch.most_common()),
        ]
        if cleared:
            lines.append(f"  old barcode cleared from {len(cleared)} other vehicle(s): {', '.join(cleared)}")
        return "\n".join(lines)
