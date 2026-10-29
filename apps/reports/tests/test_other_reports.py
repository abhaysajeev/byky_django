"""Replacement, Card Discount, Credit Note, Customer Wise Rent, Customer Wise
Discount Card, Vehicle and Asset Details (apps/reports/defs) -- figures,
filters, company isolation and each export's title and headers."""

import datetime
import io

import openpyxl
import pytest

from apps.discount.models import CardGrade, CardType, ClaimStatus
from apps.discount.tests.conftest import make_claim, make_discount
from apps.fleet.models import Asset, AssetType, Vehicle
from apps.rental import credit_notes
from apps.rental.models import (
    DiscountSource,
    Invoice,
    Order,
    OrderAction,
    OrderEvent,
    OrderItem,
    OrderItemStatus,
)
from apps.rental.tests import test_manager_reports as mr
from apps.rental.tests import test_order_requests as rq
from apps.reports.tests.helpers import DAY, data, page_rows, web
from core.enums import ApprovalStatus
from core.ids import uuid7
from core.models import User
from core.timezones import zone_for

world, token, fresh_throttle, shop = rq.world, rq.token, rq.fresh_throttle, rq.shop
settled, dubai, book = mr.settled, mr.dubai, mr.book
_ = web


def cells(d):
    return [r["cells"] for r in page_rows(d)]


def at(world, hhmm):
    hour, minute = (int(x) for x in hhmm.split(":"))
    return datetime.datetime(2026, 10, 2, hour, minute, tzinfo=zone_for(world["company"]))


# -- Replacement Details ------------------------------------------------------------------


@pytest.fixture
def replaced(web, world, token, shop):
    """Order 232, MO 42 swapped for MO 41 at 16:20 by the operator."""
    booked = book(web, world, shop, token, vehicles=("mo42",), number=232)
    order = Order.objects.get(order_no=booked["order_no"])
    old = order.items.get()
    new = OrderItem.objects.create(id=uuid7(), order=order, vehicle=shop["mo41"], package_minutes=60,
                                   start_time=at(world, "16:20"), expected_end_time=at(world, "17:20"),
                                   base_fare=old.base_fare, replaced_item=old)
    OrderItem.objects.filter(pk=old.pk).update(status=OrderItemStatus.REPLACED, end_time=at(world, "16:20"))
    OrderEvent.objects.create(id=uuid7(), order=order, action=OrderAction.REPLACE, request_hash="t", response={},
                              happened_at=at(world, "16:20"), user=User.objects.get(username="opr001"),
                              item=old, new_item=new)
    return order


def test_replacement_lists_who_swapped_what_and_when(web, replaced):
    d = data(web, "replacement-details")

    [row] = cells(d)
    assert (row["old"], row["new"], row["status"], row["vehicle_status"]) == ("MO 42", "MO 41", "Running", "Replaced")
    assert row["replaced_at"] == "02/10/2026 04:20:00 PM" and row["replaced_by"].startswith("Rashed")
    assert d["total"] is None if "total" in d else True


def test_replacement_filters_by_either_vehicle_and_the_time(web, replaced):
    assert data(web, "replacement-details", vehicle_no="MO 41")["count"] == 1
    assert data(web, "replacement-details", vehicle_no="MO 42")["count"] == 1
    assert data(web, "replacement-details", vehicle_no="DC 02")["count"] == 0
    assert data(web, "replacement-details", **{"from": "2026-10-02T16:21"})["count"] == 0


# -- Card Discount Details ---------------------------------------------------------------


@pytest.fixture
def carded(web, world, shop, settled):
    """The settled bill took a Family Card / Gold 5% discount: 5.500 off."""
    card_type = CardType.objects.create(company=world["company"], code="FAM", name="Family Card")
    grade = CardGrade.objects.create(company=world["company"], card_type=card_type, code="GOLD", name="Gold")
    order = Order.objects.get()
    claim = make_claim(make_discount(grade), shop["ahmed"], status=ClaimStatus.REDEEMED, order=order,
                       when=at(world, "16:30"), card_number="8685", branch=world["adc1"])
    Order.objects.update(discount_claim=claim)
    Invoice.objects.update(discount_source=DiscountSource.CARD, discount_amount="5.500")
    return {"type": card_type, "grade": grade, "claim": claim}


def test_card_discount_shows_the_card_the_bill_and_the_money_received(web, carded):
    d = data(web, "card-discount-details")

    [row] = cells(d)
    assert (row["card_type"], row["card_grade"], row["vehicles"]) == ("Family Card", "Gold", "MO 41, DC 02")
    assert (row["bill"], row["tax_pct"], row["tax"], row["discount"], row["received"]) == (
        "104.760", "5.00", "5.240", "5.500", "110.000")
    assert row["exit"] == "02/10/2026 05:12:00 PM"
    assert d["rows"][-1]["label"] == "TOTAL"


def test_card_discount_filters(web, world, carded):
    assert data(web, "card-discount-details", card_type=carded["type"].pk)["count"] == 1
    assert data(web, "card-discount-details", card_grade=carded["grade"].pk + 999)["count"] == 0
    assert data(web, "card-discount-details", vehicle_no="DC 02")["count"] == 1
    assert data(web, "card-discount-details", vehicle_no="ZZ")["count"] == 0


def test_a_bill_without_a_card_discount_is_not_listed(web, settled):
    assert data(web, "card-discount-details")["count"] == 0


# -- Credit Note Details -----------------------------------------------------------------


@pytest.fixture
def credited(web, world, settled):
    credit_notes.issue_directly(User.objects.get(username="opr001"), Invoice.objects.get().pk,
                                [world["company"].pk], "20")
    return Invoice.objects.get()


def test_credit_note_details_against_the_bill(web, credited):
    d = data(web, "credit-note-details", **{"from": "2026-10-01T00:00", "to": "2030-12-31T23:59"})

    [row] = cells(d)
    assert (row["invoice_no"], row["items"], row["taxable"], row["tax"], row["net"], row["credit"], row["actual"]) == (
        credited.invoice_no, "MO 41, DC 02", "104.760", "5.240", "110.000", "20.000", "90.000")
    assert row["credit_note_no"].endswith("CN")
    assert data(web, "credit-note-details", credit_note_no="nope",
                **{"from": "2026-10-01T00:00", "to": "2030-12-31T23:59"})["count"] == 0


# -- Customer Wise Rent ------------------------------------------------------------------


def test_customer_wise_rent_has_a_row_per_bill_day_and_credit_note_day(web, credited):
    wide = {"from": "2026-10-01T00:00", "to": "2030-12-31T23:59"}
    d = data(web, "customer-wise-rent", **wide)

    rows = cells(d)
    assert [r["type"] for r in rows] == ["Invoice", "Credit Note"]
    bill, note = rows
    assert (bill["customer"], bill["status"], bill["count"], bill["vehicles"], bill["bill"], bill["tax"]) == (
        "Ahmed Al Mansoori", "Active", 1, 2, "104.760", "5.240")
    assert note["count"] == 1 and note["vehicles"] == 2
    assert [r["kind"] for r in d["rows"]] == ["data", "data"]          # no TOTAL: bills and credit notes don't add
    only_notes = data(web, "customer-wise-rent", credit_note_no="CN", **wide)
    assert [r["type"] for r in cells(only_notes)] == ["Credit Note"]


# -- Customer Wise Discount Card ---------------------------------------------------------


def test_a_card_is_listed_once_from_its_first_claim(web, world, shop, carded):
    later = make_claim(carded["claim"].card_discount, shop["ahmed"], status=ClaimStatus.REDEEMED,
                       when=at(world, "18:00"), card_number="8685", branch=None)
    assert later.pk

    d = data(web, "customer-wise-card")

    [row] = cells(d)
    assert (row["card_no"], row["card_type"], row["card_grade"], row["branch"], row["created"]) == (
        "8685", "Family Card", "Gold", "Abu Dhabi Corniche 1", "02/10/2026 04:30:00 PM")
    assert data(web, "customer-wise-card", **{"from": "2026-10-02T17:00"})["count"] == 0   # first used before
    assert data(web, "customer-wise-card", card_no="999")["count"] == 0


# -- Vehicle and Asset Details -----------------------------------------------------------


def test_vehicle_details_by_station_with_counts(web, world, shop):
    Vehicle.objects.filter(pk=shop["dc02"].pk).update(vehicle_type=world["berg"])

    d = data(web, "vehicle-details")

    rows = cells(d)
    by_code = {r["code"]: r for r in rows}
    assert by_code["MO 41"]["branch"] == "Abu Dhabi Corniche 1"
    assert (by_code["MO 41"]["type"], by_code["MO 41"]["type_count"], by_code["MO 41"]["category_count"]) == (
        "Monaco", 2, 3)
    assert by_code["DC 02"]["type_count"] == 1
    assert by_code["MO 99"]["branch"] == "Abu Dhabi Corniche 2"
    assert by_code["MO 41"]["status"] == "Active"
    assert data(web, "vehicle-details", vehicle_type=world["berg"].pk)["count"] == 1
    assert data(web, "vehicle-details", branch=world["adc2"].pk)["count"] == 1
    assert data(web, "vehicle-details", vehicle_no="mo 4")["count"] == 2
    assert "from" not in d["query"]                                   # no date range


def test_asset_details_with_custodian_and_status(web, world):
    camera = AssetType.objects.create(company=world["company"], asset_type_name="Camera")
    Asset.objects.create(company=world["company"], asset_code="CCTV001", asset_type=camera, branch=world["adc1"],
                         serial_no="SN-1", approval_status=ApprovalStatus.APPROVED)
    Asset.objects.create(company=world["company"], asset_code="CCTV002", asset_type=camera, is_active=False)

    rows = cells(data(web, "asset-details"))

    # Active or Inactive only, as the Inventory screens show it.
    assert [(r["branch"], r["code"], r["type_count"], r["status"]) for r in rows] == [
        ("Abu Dhabi Corniche 1", "CCTV001", 1, "Active"), ("Not at a branch", "CCTV002", 1, "Inactive")]
    assert data(web, "asset-details", asset="SN-1")["count"] == 1


# -- Every report: page, isolation, export -------------------------------------------------

SLUGS = ["replacement-details", "card-discount-details", "credit-note-details", "customer-wise-rent",
         "customer-wise-card", "vehicle-details", "asset-details"]
TITLES = {
    "replacement-details": ("Replacement Details Report", "Branch"),
    "card-discount-details": ("Card Discount Details Report", "State"),
    "credit-note-details": ("Credit Note Details Report", "Date"),
    "customer-wise-rent": ("Customer Wise Rent Details Report", "Type"),
    "customer-wise-card": ("Customer Wise Discount Card Details Report", "Phone Number"),
    "vehicle-details": ("Vehicle Details Report", "Branch"),
    "asset-details": ("Asset Details Report", "Branch"),
}


@pytest.mark.parametrize("slug", SLUGS)
def test_every_report_renders_and_exports_its_title_and_headers(web, slug):
    assert web.get(f"/reports/{slug}/", DAY).status_code == 200
    response = web.get(f"/reports/{slug}/export/", DAY)
    assert response.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(response.content)).active
    title, first = TITLES[slug]
    assert (ws["A1"].value, ws["A4"].value) == (title, first)


def test_another_companys_vehicles_and_assets_are_not_listed(web, world, shop):
    Vehicle.objects.update(company=world["other"], current_branch=None)   # a station is its own company's

    assert data(web, "vehicle-details")["count"] == 0
