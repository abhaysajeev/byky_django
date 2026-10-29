"""The money reports (apps/reports/defs): Rent Summary, Consolidated and Rent
Details -- their figures on real bookings through the tablet API, every
filter, paging, the permissions, and the Excel export's layout.

The settled order (apps/rental/tests/test_manager_reports.py `settled`): MO 41
and DC 02 booked 2 Oct 16:00 at Abu Dhabi Corniche 1, 100.000 cash down; MO 41
back twelve minutes late at 60.000, DC 02 on time at 50.000; settled 17:13 at
net 110.000 with 5.240 VAT inside, the last 10.000 by card (reference 4421).
"""

import io

import openpyxl
import pytest

from apps.company.models import PaymentMode
from apps.fleet.models import Vehicle
from apps.portal.models import RolePermission
from apps.rental import credit_notes
from apps.rental.models import DiscountSource, Invoice, Order, OrderStatus
from apps.rental.tests import test_manager_reports as mr
from apps.rental.tests import test_order_requests as rq
from apps.reports import views
from apps.reports.common import bucket
from apps.reports.tests.helpers import DAY, data, page_rows, web
from core.models import User

world, token, fresh_throttle, shop = rq.world, rq.token, rq.fresh_throttle, rq.shop
settled, dubai, book = mr.settled, mr.dubai, mr.book
_ = web     # a fixture


def cells(d):
    return [r["cells"] for r in page_rows(d)]


# -- Payment modes -------------------------------------------------------------------


@pytest.mark.parametrize(("name", "key"), [
    ("Cash", "cash"), ("Card", "card"), ("Visa Card", "card"), ("Tap & Pay", "tap"), ("Tap and pay card", "tap"),
    ("UPI", "upi"), ("Cheque", "other")])
def test_payment_modes_fall_in_their_column_by_name(name, key):
    assert bucket(name) == key


# -- Rent Summary ---------------------------------------------------------------------


def test_rent_summary_is_settled_orders_per_day_and_station_by_state(web, settled):
    d = data(web, "rent-summary")

    assert [r["kind"] for r in d["rows"]] == ["data", "total", "grand"]
    [row] = cells(d)
    assert row == {"date": "02/10/2026", "state": "Abu Dhabi", "branch": "Abu Dhabi Corniche 1", "orders": 1,
                   "net": "110.000", "credit": "0.000", "actual": "110.000", "pay_cash": "100.000",
                   "pay_card": "10.000", "pay_tap": "0.000", "pay_upi": "0.000"}
    assert d["rows"][2]["cells"]["net"] == "110.000" and d["rows"][2]["label"] == "GRAND TOTAL"
    assert [k["value"] for k in d["kpis"]] == [1, "110.000", "0.000", "110.000"]
    assert "pay_other" not in [c["key"] for c in d["columns"]]


def test_rent_summary_takes_approved_credit_notes_off_the_actual_collection(web, world, settled):
    credit_notes.issue_directly(User.objects.get(username="opr001"), Invoice.objects.get().pk,
                                [world["company"].pk], "20")

    [row] = cells(data(web, "rent-summary"))

    assert (row["net"], row["credit"], row["actual"]) == ("110.000", "20.000", "90.000")


def test_rent_summary_follows_the_settle_time_and_the_place(web, settled, dubai):
    assert data(web, "rent-summary", **{"from": "2026-10-02T17:14", "to": "2026-10-02T23:59"})["count"] == 0
    assert data(web, "rent-summary", **{"from": "2026-10-02T17:13", "to": "2026-10-02T17:13"})["count"] == 1
    assert data(web, "rent-summary", state=dubai["state"].pk)["count"] == 0
    assert data(web, "rent-summary", branch=dubai["open"].pk)["count"] == 0


def test_a_payment_mode_that_is_none_of_the_four_gets_an_other_column(web, world, settled):
    cheque = PaymentMode.objects.create(company=world["company"], name="Cheque")
    Order.objects.get().payments.filter(mode__name="Card").update(mode=cheque)

    d = data(web, "rent-summary")

    assert "pay_other" in [c["key"] for c in d["columns"]]
    [row] = cells(d)
    assert (row["pay_card"], row["pay_other"]) == ("0.000", "10.000")


# -- Consolidated ---------------------------------------------------------------------


@pytest.fixture
def mixed(web, world, token, shop, settled):
    """Beside the settled 231: 232 still running (MO 42), 233 cancelled."""
    book(web, world, shop, token, vehicles=("mo42",), number=232)
    cancelled = book(web, world, shop, token, vehicles=("mo41",), number=233)
    Order.objects.filter(order_no=cancelled["order_no"]).update(status=OrderStatus.CANCELLED)
    return settled


def test_consolidated_is_one_row_per_day_station_and_tablet(web, mixed):
    d = data(web, "consolidated")

    [row] = cells(d)
    assert (row["state"], row["branch"], row["orders"], row["running"], row["cancelled"]) == (
        "Abu Dhabi", "Abu Dhabi Corniche 1", 3, 1, 1)
    assert (row["first_no"], row["last_no"]) == ("231", "233")
    assert (row["bill"], row["tax"], row["net"], row["pay_cash"], row["pay_card"], row["actual"]) == (
        "104.760", "5.240", "110.000", "100.000", "10.000", "110.000")
    assert d["rows"][-1]["label"] == "TOTAL" and d["rows"][-1]["cells"]["orders"] == 3


def test_consolidated_counts_discounted_and_complimentary_bills(web, settled):
    Invoice.objects.update(discount_source=DiscountSource.COMPLIMENTARY)
    [row] = cells(data(web, "consolidated"))
    assert (row["complimentary"], row["discounted"]) == (1, 0)

    Invoice.objects.update(discount_source=DiscountSource.CARD)
    [row] = cells(data(web, "consolidated"))
    assert (row["complimentary"], row["discounted"]) == (0, 1)


def test_consolidated_filters_by_device_and_cashier(web, mixed):
    operator = User.objects.get(username="opr001")

    assert data(web, "consolidated", device="no such tablet")["count"] == 0
    assert data(web, "consolidated", cashier=operator.pk)["count"] == 1
    assert data(web, "consolidated", cashier=User.objects.get(username="rep.user").pk)["count"] == 0


# -- Rent Details ---------------------------------------------------------------------


def test_rent_details_shows_the_bill_once_and_each_vehicle(web, settled):
    d = data(web, "rent-details")

    [order] = page_rows(d)
    c = order["cells"]
    assert (c["status"], c["customer"], c["initial"], c["collected"], c["deposit"]) == (
        "Received", "Ahmed Al Mansoori", "100.000", "110.000", "0.000")
    assert (c["tax_pct"], c["tax"], c["net"], c["credit"], c["actual"], c["bill"]) == (
        "5.00", "5.240", "110.000", "0.000", "110.000", "104.760")
    assert (c["pay_cash"], c["pay_card"], c["refs"]) == ("100.000", "10.000", "4421")
    assert (c["supervisor"], c["card_discount"], c["card_type"]) == ("0.000", "0.000", None)
    assert c["open_by"] == c["close_by"] and c["open_by"].startswith("Rashed")
    assert [(ln["vehicle_no"], ln["vehicle_status"], ln["initial_fare"], ln["total_fare"]) for ln in order["lines"]] == [
        ("MO 41", "Received", "50.000", "60.000"), ("DC 02", "Received", "50.000", "50.000")]
    total = d["rows"][-1]
    assert (total["label"], total["cells"]["initial_fare"], total["cells"]["total_fare"], total["cells"]["net"]) == (
        "TOTAL", "100.000", "110.000", "110.000")
    assert d["note"] == ""


def test_a_vehicle_filter_shows_matching_lines_only_with_the_whole_bill(web, world, shop, settled):
    Vehicle.objects.filter(pk=shop["dc02"].pk).update(vehicle_type=world["berg"])

    d = data(web, "rent-details", vehicle_type=world["berg"].pk)

    [order] = page_rows(d)
    assert [ln["vehicle_no"] for ln in order["lines"]] == ["DC 02"]
    assert (order["cells"]["net"], order["cells"]["initial"]) == ("110.000", "100.000")
    total = d["rows"][-1]["cells"]
    assert (total["total_fare"], total["net"]) == ("50.000", "110.000")
    assert d["note"] == views_note()
    assert data(web, "rent-details", vehicle_no="MO 4")["count"] == 1
    assert data(web, "rent-details", vehicle_no="ZZ")["count"] == 0


def views_note():
    from apps.reports.defs.rent_details import VEHICLE_NOTE
    return VEHICLE_NOTE


@pytest.mark.parametrize(("status", "found"), [
    ("received", ["231"]), ("running", ["232"]), ("cancelled", ["233"]), ("partially_received", []),
    ("complimentary", []), ("direct_rent", []), ("discounted", [])])
def test_rent_details_status_filter(web, mixed, status, found):
    d = data(web, "rent-details", status=status)

    assert [c["order_no"][-3:] for c in cells(d)] == found


def test_a_complimentary_bill_is_a_supervisor_discount(web, settled):
    Invoice.objects.update(discount_source=DiscountSource.COMPLIMENTARY, discount_amount="110.000")

    [c] = cells(data(web, "rent-details", status="complimentary"))

    assert (c["supervisor"], c["card_discount"]) == ("110.000", "0.000")


def test_rent_details_finds_by_customer_phone_and_order_no(web, settled):
    assert data(web, "rent-details", customer="ahmed")["count"] == 1
    assert data(web, "rent-details", customer="nobody")["count"] == 0
    assert data(web, "rent-details", phone="501234567")["count"] == 1
    assert data(web, "rent-details", order_no="231")["count"] == 1


def test_rent_details_pages_by_order_and_totals_every_page(web, mixed, monkeypatch):
    monkeypatch.setattr(views, "PER_PAGE", 1)

    first = data(web, "rent-details")
    last = data(web, "rent-details", page=3)

    assert (first["count"], first["pages"], len(page_rows(first))) == (3, 3, 1)
    assert [r["kind"] for r in first["rows"]] == ["data"]           # the TOTAL comes on the last page
    assert last["rows"][-1]["label"] == "TOTAL" and last["rows"][-1]["cells"]["net"] == "110.000"
    assert first["kpis"][2]["value"] == "110.000"                    # the headline figures are every order's


# -- Permissions --------------------------------------------------------------------------


def test_the_page_renders_with_its_first_page(web, settled):
    page = web.get("/reports/rent-details/", DAY)

    assert page.status_code == 200
    assert b'id="rpt-initial"' in page.content and b"Rent Details" in page.content


def test_read_is_needed_to_see_and_print_to_export(web):
    RolePermission.objects.filter(role=web.role, page__code="reports.rent_summary").update(can_print=False)
    assert web.get("/reports/rent-summary/export/").status_code == 403
    assert web.get("/reports/rent-summary/").status_code == 200

    RolePermission.objects.filter(role=web.role, page__code="reports.rent_summary").update(can_read=False)
    assert web.get("/reports/rent-summary/").status_code == 403
    assert web.get("/reports/rent-summary/data/").status_code == 403


def test_signed_out_the_data_says_so(client, db):
    assert client.get("/reports/rent-summary/data/").status_code == 401


def test_another_companys_orders_are_not_in_the_reports(web, world, settled):
    Order.objects.update(company=world["other"])
    Invoice.objects.update(company=world["other"])

    assert data(web, "rent-details")["count"] == 0
    assert data(web, "rent-summary")["count"] == 0


# -- Excel ----------------------------------------------------------------------------------


def sheet(response):
    assert response.status_code == 200, response.content[:300]
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    return openpyxl.load_workbook(io.BytesIO(response.content)).active


def test_rent_details_exports_like_the_sample(web, settled):
    ws = sheet(web.get("/reports/rent-details/export/", DAY))

    assert ws["A1"].value == "Rent Details Report"
    assert ws["A2"].value.startswith("From : 02/10/2026 12:00 AM")
    assert ws["A3"].value.startswith("Report Generated On : ")
    headers = [c.value for c in ws[4]]
    assert headers[:4] == ["Branch", "Order Number", "Order Date And Time", "Order Status"]
    assert "Vehicle Number" in headers and "SupervisorDiscount Amount (AED)" in headers
    assert ws["B5"].value.endswith("231") and ws["B6"].value is None
    assert "B5:B6" in {str(r) for r in ws.merged_cells.ranges}           # the order, merged down its vehicles
    vehicle = headers.index("Vehicle Number") + 1
    assert [ws.cell(5, vehicle).value, ws.cell(6, vehicle).value] == ["MO 41", "DC 02"]
    net = headers.index("Net Amount (AED)") + 1
    assert float(ws.cell(5, net).value) == 110.0 and ws.cell(5, net).number_format == "#,##0.000"
    assert ws.cell(7, 1).value == "TOTAL"


def test_rent_summary_exports_its_blocks_and_grand_total(web, settled):
    ws = sheet(web.get("/reports/rent-summary/export/", DAY))

    assert ws["A1"].value == "Rent Summary Report"
    assert [c.value for c in ws[4]][:5] == ["Date", "State", "Branch", "Total No Of Order", "Net Bill Amount (AED)"]
    assert (ws["C5"].value, ws["A6"].value, ws["A7"].value) == ("Abu Dhabi Corniche 1", "TOTAL", "GRAND TOTAL")
    assert ws["A5"].number_format == "dd/mm/yyyy"


def test_an_export_past_the_limit_asks_for_narrower_filters(web, settled, monkeypatch):
    monkeypatch.setattr(views, "EXPORT_LIMIT", 0)

    response = web.get("/reports/rent-details/export/", DAY)

    assert response.status_code == 413 and "Narrow the dates" in response.json()["detail"]
